"""Generate a styled weekly Excel report from the synthetic ETL outputs."""

from __future__ import annotations

import argparse
import shutil
from datetime import datetime
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="12355B")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _append_table(sheet, headers: list[str], rows: list[list[object]]) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, column in enumerate(sheet.columns, start=1):
        width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_weekly_report(incidents_csv: Path, kpis_csv: Path, output: Path) -> Path:
    import csv

    with incidents_csv.open(encoding="utf-8-sig", newline="") as handle:
        incident_reader = csv.DictReader(handle)
        incident_rows = list(incident_reader)
        incident_headers = incident_reader.fieldnames or []
    with kpis_csv.open(encoding="utf-8-sig", newline="") as handle:
        kpi_reader = csv.DictReader(handle)
        kpi_rows = list(kpi_reader)
        kpi_headers = kpi_reader.fieldnames or []

    if output.exists():
        backup_dir = output.parent / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(output, backup_dir / f"{output.stem}_{timestamp}{output.suffix}")

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary.append(["CoreOps - relatório semanal sintético"])
    summary["A1"].font = Font(size=16, bold=True, color="12355B")
    summary.append(["Indicador", "Valor"])
    for cell in summary[2]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    summary.append(["Incidentes processados", len(incident_rows)])
    summary.append(["Incidentes ativos", sum(row.get("is_active") == "True" for row in incident_rows)])
    summary.append(["Operadoras monitoradas", len(kpi_rows)])
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 18

    incidents_sheet = workbook.create_sheet("Incidentes")
    _append_table(
        incidents_sheet,
        incident_headers,
        [[row.get(header, "") for header in incident_headers] for row in incident_rows],
    )
    kpis_sheet = workbook.create_sheet("KPIs por operadora")
    _append_table(
        kpis_sheet,
        kpi_headers,
        [[row.get(header, "") for header in kpi_headers] for row in kpi_rows],
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Gera relatório Excel semanal a partir do ETL sintético.")
    parser.add_argument("--incidents", type=Path, default=Path("data/processed/incidents.csv"))
    parser.add_argument("--kpis", type=Path, default=Path("data/processed/provider_kpis.csv"))
    parser.add_argument("--output", type=Path, default=Path("output/coreops_weekly_demo.xlsx"))
    return parser.parse_args(argv)


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(argv)
    print(build_weekly_report(args.incidents, args.kpis, args.output))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
