from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from coreops_demo.analysis import build_analysis
from coreops_demo.pipeline import (
    DEFAULT_AS_OF,
    deduplicate_events,
    enrich_events,
    filter_operational_events,
    load_sources,
    normalize_key,
    run_pipeline,
)
from coreops_demo.weekly_report import build_weekly_report


ROOT = Path(__file__).resolve().parents[1]


def test_normalize_key_handles_accents_and_punctuation():
    assert normalize_key("  Clínica São José / WAN-1 ") == "CLINICASAOJOSEWAN1"


def test_filter_deduplicate_and_fallback_matching():
    events, inventory = load_sources(ROOT / "data" / "raw")
    operational = filter_operational_events(events)
    deduplicated = deduplicate_events(operational)
    enriched = enrich_events(deduplicated, inventory)
    assert len(events) > len(operational)
    assert len(operational) > len(deduplicated)
    assert "circuit_fallback" in set(enriched["match_method"])
    assert enriched["unit"].notna().all()


def test_pipeline_generates_dashboard_and_quality_artifacts(tmp_path):
    artifacts = run_pipeline(ROOT / "data" / "raw", tmp_path, DEFAULT_AS_OF)
    incidents = pd.read_csv(artifacts.incidents)
    kpis = pd.read_csv(artifacts.provider_kpis)
    assert not incidents.empty
    assert incidents["is_active"].astype(str).eq("True").any()
    assert set(kpis.columns) >= {"operator", "mttr_minutes", "availability_pct"}
    assert "window.COREOPS_DEMO_DATA" in artifacts.dashboard_javascript.read_text(encoding="utf-8")


def test_weekly_report_contains_summary_incidents_and_kpis(tmp_path):
    artifacts = run_pipeline(ROOT / "data" / "raw", tmp_path / "processed", DEFAULT_AS_OF)
    report = build_weekly_report(artifacts.incidents, artifacts.provider_kpis, tmp_path / "weekly.xlsx")
    workbook = load_workbook(report, read_only=True)
    assert workbook.sheetnames == ["Resumo", "Incidentes", "KPIs por operadora"]


def test_analysis_generates_statistics_recommendations_and_charts(tmp_path):
    pipeline = run_pipeline(ROOT / "data" / "raw", tmp_path / "processed", DEFAULT_AS_OF)
    artifacts = build_analysis(pipeline.incidents, tmp_path / "analysis")
    providers = pd.read_csv(artifacts.provider_summary)
    correlations = pd.read_csv(artifacts.correlation_summary)
    report = artifacts.report.read_text(encoding="utf-8")

    assert set(providers.columns) >= {"operator", "incident_count", "downtime_share_pct"}
    assert len(correlations) == 2
    assert "## Perguntas de negócio" in report
    assert "## Recomendações" in report
    assert "## Limitações" in report
    assert artifacts.provider_chart.read_text(encoding="utf-8").startswith("<svg")
    assert artifacts.monthly_chart.read_text(encoding="utf-8").startswith("<svg")
