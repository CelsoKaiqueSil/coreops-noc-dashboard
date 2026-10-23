"""Check actual exported XLSX files and ZIP integrity."""
from pathlib import Path
import zipfile
from openpyxl import load_workbook

paths = list(Path('output').glob('*.xlsx'))
assert paths, 'No synthetic workbook generated'
for path in paths:
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, path
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        assert book.sheetnames
        for sheet in book:
            for row in sheet.iter_rows():
                for cell in row:
                    assert cell.data_type != 'e', (path.name, sheet.title, cell.coordinate)
    finally:
        book.close()
print(f'{len(paths)} workbook(s) validated')
